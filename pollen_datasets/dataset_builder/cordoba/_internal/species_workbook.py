"""Read and validate the species workbook for preparation and inference."""

from pathlib import Path

from openpyxl import load_workbook


MONTHS = (
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
)


def normalize_species(name):
    return name.strip().casefold()


def read_species_workbook(path, *, monthly=False, sheet_name="Monthly labels"):
    """Return normalized flags while preserving species spelling.

    Exclusions require species/known only; monthly export also requires month/active.
    Own the input stream so validation failures release Windows file handles.
    """
    required = ("month", "species", "active", "known") if monthly else ("species", "known")
    records = []
    known_flags = {}
    month_names = {month.casefold(): month for month in MONTHS}
    with Path(path).open("rb") as stream:
        book = load_workbook(stream, read_only=True, data_only=True)
        try:
            if sheet_name not in book.sheetnames:
                raise ValueError(f"Worksheet {sheet_name!r} was not found.")
            rows = iter(book[sheet_name].values)
            headers = [str(value).strip().casefold() for value in next(rows, ())]
            for name in required:
                if headers.count(name) != 1:
                    raise ValueError(f"Expected exactly one {name!r} column in {sheet_name!r}.")
            positions = {name: headers.index(name) for name in required}
            for number, row in enumerate(rows, 2):
                if all(value is None for value in row):
                    continue
                record = {name: row[index] if index < len(row) else None
                          for name, index in positions.items()}
                species = record["species"]
                if not isinstance(species, str) or not species.strip():
                    raise ValueError(f"Row {number}: species must be a nonempty name.")
                record["species"] = species.strip()
                for name in (("active", "known") if monthly else ("known",)):
                    flag = str(record[name]).strip().casefold()
                    if flag not in ("yes", "no"):
                        raise ValueError(f"Row {number}: {name} must be yes/no, got {record[name]!r}.")
                    record[name] = flag == "yes"
                key = normalize_species(species)
                if key in known_flags and known_flags[key] != record["known"]:
                    raise ValueError(f"Inconsistent known flags for {species.strip()}")
                known_flags[key] = record["known"]
                if monthly:
                    month = month_names.get(str(record["month"]).strip().casefold())
                    if month is None:
                        raise ValueError(f"Row {number}: invalid month {record['month']!r}.")
                    record["month"] = month
                records.append(record)
        finally:
            book.close()
    return records
